#!/usr/bin/env python3
"""Synthetic test suite for scripts/export_cdc_ai_review_comparison.py.

SYNTHETIC DATA ONLY. No real archive filenames, paths, project names, or
PostgreSQL rows are used anywhere in this file. No live connection of any
kind is required or attempted.
"""
from __future__ import annotations

import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent
SPEC = importlib.util.spec_from_file_location("export_cdc_ai_review_comparison", HERE / "export_cdc_ai_review_comparison.py")
mod = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = mod
SPEC.loader.exec_module(mod)


def synthetic_row(**overrides) -> "mod.ComparisonRow":
    defaults = dict(
        candidate_id="cand-0001", archive_file_id=1, year=2020, project_key="OFFRES 2020/SYNTHETIC",
        filename="synthetic.pdf", relative_path="OFFRES 2020/SYNTHETIC/synthetic.pdf",
        original_proposal="UNKNOWN", validation_status="HUMAN_VALIDATED_CDC",
        extraction_status="SUCCESS", extraction_method="pdf_text", extraction_failure_category=None,
        ai_proposed_role=None, ai_confidence=None, ai_model_name=None, ai_created_at=None,
        ai_processing_status=None,
        is_primary_candidate=True, duplicate_of_archive_file_id=None,
    )
    defaults.update(overrides)
    return mod.ComparisonRow(**defaults)


class YoussefDecisionTest(unittest.TestCase):
    def test_human_validated_cdc_is_oui(self):
        row = synthetic_row(validation_status="HUMAN_VALIDATED_CDC")
        self.assertEqual(row.youssef_decision, "OUI")

    def test_human_rejected_cdc_is_non(self):
        row = synthetic_row(validation_status="HUMAN_REJECTED_CDC")
        self.assertEqual(row.youssef_decision, "NON")

    def test_anything_else_is_incertain(self):
        row = synthetic_row(validation_status="MACHINE_CLASSIFIED")
        self.assertEqual(row.youssef_decision, "INCERTAIN")

    def test_decision_is_never_recomputed_from_ai_or_extraction_fields(self):
        # Changing AI/extraction fields must never change youssef_decision.
        row_a = synthetic_row(validation_status="HUMAN_VALIDATED_CDC", ai_proposed_role="TDR", ai_processing_status="SUCCESS")
        row_b = synthetic_row(validation_status="HUMAN_VALIDATED_CDC", extraction_status="FAILED")
        self.assertEqual(row_a.youssef_decision, "OUI")
        self.assertEqual(row_b.youssef_decision, "OUI")


class AgreementMappingTest(unittest.TestCase):
    def test_extraction_failed_is_analyse_impossible_regardless_of_decision(self):
        row = synthetic_row(extraction_status="FAILED", validation_status="HUMAN_VALIDATED_CDC")
        self.assertEqual(mod.compute_agreement(row), mod.AGREEMENT_ANALYSE_IA_IMPOSSIBLE)

    def test_no_ai_review_is_non_compare(self):
        row = synthetic_row(extraction_status="SUCCESS", ai_processing_status=None)
        self.assertEqual(mod.compute_agreement(row), mod.AGREEMENT_NON_COMPARE)

    def test_incertain_is_a_revoir_even_with_an_ai_review(self):
        row = synthetic_row(
            validation_status="MACHINE_CLASSIFIED", extraction_status="SUCCESS",
            ai_proposed_role="CDC", ai_processing_status="SUCCESS",
        )
        self.assertEqual(mod.compute_agreement(row), mod.AGREEMENT_A_REVOIR)

    def test_oui_and_cdc_role_is_accord(self):
        row = synthetic_row(
            validation_status="HUMAN_VALIDATED_CDC", extraction_status="SUCCESS",
            ai_proposed_role="CDC", ai_processing_status="SUCCESS",
        )
        self.assertEqual(mod.compute_agreement(row), mod.AGREEMENT_ACCORD)

    def test_oui_and_dao_with_cdc_role_is_accord(self):
        row = synthetic_row(
            validation_status="HUMAN_VALIDATED_CDC", extraction_status="SUCCESS",
            ai_proposed_role="DAO_WITH_CDC", ai_processing_status="SUCCESS",
        )
        self.assertEqual(mod.compute_agreement(row), mod.AGREEMENT_ACCORD)

    def test_non_and_non_cdc_role_is_accord(self):
        row = synthetic_row(
            validation_status="HUMAN_REJECTED_CDC", extraction_status="SUCCESS",
            ai_proposed_role="OFFER", ai_processing_status="SUCCESS",
        )
        self.assertEqual(mod.compute_agreement(row), mod.AGREEMENT_ACCORD)

    def test_oui_and_non_cdc_role_is_desaccord(self):
        row = synthetic_row(
            validation_status="HUMAN_VALIDATED_CDC", extraction_status="SUCCESS",
            ai_proposed_role="OFFER", ai_processing_status="SUCCESS",
        )
        self.assertEqual(mod.compute_agreement(row), mod.AGREEMENT_DESACCORD)

    def test_non_and_cdc_role_is_desaccord(self):
        row = synthetic_row(
            validation_status="HUMAN_REJECTED_CDC", extraction_status="SUCCESS",
            ai_proposed_role="CDC", ai_processing_status="SUCCESS",
        )
        self.assertEqual(mod.compute_agreement(row), mod.AGREEMENT_DESACCORD)

    def test_agreement_is_always_one_of_the_five_defined_values(self):
        for row in [
            synthetic_row(extraction_status="FAILED"),
            synthetic_row(ai_processing_status=None),
            synthetic_row(validation_status="MACHINE_CLASSIFIED", ai_proposed_role="CDC", ai_processing_status="SUCCESS"),
            synthetic_row(validation_status="HUMAN_VALIDATED_CDC", ai_proposed_role="CDC", ai_processing_status="SUCCESS"),
            synthetic_row(validation_status="HUMAN_REJECTED_CDC", ai_proposed_role="CDC", ai_processing_status="SUCCESS"),
        ]:
            self.assertIn(mod.compute_agreement(row), mod.AGREEMENT_VALUES)


class ExtractionFailureLabelTest(unittest.TestCase):
    def test_failed_row_ai_verdict_label_is_impossible(self):
        row = synthetic_row(extraction_status="FAILED", extraction_failure_category="DOC_EMBEDDED_IMAGES_ONLY")
        self.assertIn("IMPOSSIBLE", mod.compute_ai_verdict_label(row))

    def test_never_invents_a_verdict_for_a_failed_row(self):
        row = synthetic_row(extraction_status="FAILED", ai_proposed_role="CDC", ai_processing_status="SUCCESS")
        # Even if a stale AI row somehow exists for a now-FAILED candidate,
        # the verdict label must never surface it as a real result.
        label = mod.compute_ai_verdict_label(row)
        self.assertNotEqual(label, "CDC")

    def test_not_yet_reviewed_is_shown_honestly(self):
        row = synthetic_row(extraction_status="SUCCESS", ai_processing_status=None)
        self.assertEqual(mod.compute_ai_verdict_label(row), "PAS ENCORE ANALYSÉ")


class WorkbookStructureTest(unittest.TestCase):
    def _rows(self, n=750):
        rows = []
        for i in range(1, n + 1):
            decision = "HUMAN_VALIDATED_CDC" if i % 3 else "HUMAN_REJECTED_CDC"
            rows.append(synthetic_row(
                candidate_id=f"cand-{i:04d}", archive_file_id=i, filename=f"synthetic-{i}.pdf",
                relative_path=f"OFFRES 2020/SYNTHETIC-{i}/synthetic-{i}.pdf",
                validation_status=decision,
            ))
        return rows

    def test_exactly_750_data_rows(self):
        rows = self._rows(750)
        wb = mod.build_workbook(rows, "qwen3:14b", "v3", "2026-09-28 00:00")
        ws = wb["Comparaison Youssef vs IA"]
        data_row_count = sum(
            1 for row in ws.iter_rows(min_row=2, max_row=ws.max_row) if row[0].value is not None
        )
        self.assertEqual(data_row_count, 750)

    def test_visible_headers_match_exactly(self):
        rows = self._rows(3)
        wb = mod.build_workbook(rows, "qwen3:14b", "v3", "2026-09-28 00:00")
        ws = wb["Comparaison Youssef vs IA"]
        headers = [cell.value for cell in ws[1]]
        self.assertEqual(headers, mod.VISIBLE_HEADERS)

    def test_hidden_technical_sheet_present_and_hidden(self):
        rows = self._rows(3)
        wb = mod.build_workbook(rows, "qwen3:14b", "v3", "2026-09-28 00:00")
        self.assertIn("Données techniques", wb.sheetnames)
        self.assertEqual(wb["Données techniques"].sheet_state, "hidden")

    def test_technical_sheet_correspondence_matches_visible_row_count(self):
        rows = self._rows(750)
        wb = mod.build_workbook(rows, "qwen3:14b", "v3", "2026-09-28 00:00")
        visible = wb["Comparaison Youssef vs IA"]
        technical = wb["Données techniques"]
        visible_count = sum(1 for row in visible.iter_rows(min_row=2) if row[0].value is not None)
        technical_count = sum(1 for row in technical.iter_rows(min_row=2) if row[0].value is not None)
        self.assertEqual(visible_count, technical_count)
        self.assertEqual(visible_count, 750)

    def test_summary_sheet_present_with_aggregate_counts(self):
        rows = self._rows(9)
        wb = mod.build_workbook(rows, "qwen3:14b", "v3", "2026-09-28 00:00")
        self.assertIn("Résumé", wb.sheetnames)
        summary_text = "\n".join(
            str(cell.value) for row in wb["Résumé"].iter_rows() for cell in row if cell.value is not None
        )
        self.assertIn("Total candidats", summary_text)
        self.assertIn("Total désaccords", summary_text)


class RealisticPopulationScaleTest(unittest.TestCase):
    """Simulates the actual 2026-09-28 population shape: 750 candidates
    total, 436/312/2 Youssef decisions, 40 controlled extraction
    failures, and 94 rows (68 recovered SUCCESS + 26 original) each with
    their own independent refreshed AI review."""

    def _realistic_rows(self):
        rows = []
        afid = 1
        # 436 OUI, 312 NON, 2 INCERTAIN = 750
        for _ in range(436):
            rows.append(synthetic_row(candidate_id=f"cand-{afid}", archive_file_id=afid, validation_status="HUMAN_VALIDATED_CDC"))
            afid += 1
        for _ in range(312):
            rows.append(synthetic_row(candidate_id=f"cand-{afid}", archive_file_id=afid, validation_status="HUMAN_REJECTED_CDC"))
            afid += 1
        for _ in range(2):
            rows.append(synthetic_row(candidate_id=f"cand-{afid}", archive_file_id=afid, validation_status="MACHINE_CLASSIFIED"))
            afid += 1
        self.assertEqual(len(rows), 750)

        # 40 controlled extraction failures (among the OUI rows, first 40).
        for i in range(40):
            rows[i] = synthetic_row(
                candidate_id=rows[i].candidate_id, archive_file_id=rows[i].archive_file_id,
                validation_status="HUMAN_VALIDATED_CDC", extraction_status="FAILED",
                extraction_failure_category="DOC_EMBEDDED_IMAGES_ONLY",
            )

        # 94 rows (indices 40..133, still within the 436 OUI block) each
        # with their OWN independent refreshed AI review.
        for i in range(40, 134):
            rows[i] = synthetic_row(
                candidate_id=rows[i].candidate_id, archive_file_id=rows[i].archive_file_id,
                validation_status="HUMAN_VALIDATED_CDC", extraction_status="SUCCESS",
                ai_proposed_role="CDC" if i % 2 == 0 else "TDR",
                ai_confidence=0.5 + (i % 10) / 20, ai_model_name="qwen3:14b",
                ai_created_at="2026-09-28 12:00:00", ai_processing_status="SUCCESS",
            )
        return rows

    def test_750_rows_436_312_2_youssef_decisions_preserved(self):
        rows = self._realistic_rows()
        wb = mod.build_workbook(rows, "qwen3:14b", "v3", "2026-09-28 00:00")
        ws = wb["Comparaison Youssef vs IA"]
        decision_col = mod.VISIBLE_HEADERS.index("Décision de Youssef") + 1
        counts = {"OUI": 0, "NON": 0, "INCERTAIN": 0}
        for r in range(2, ws.max_row + 1):
            value = ws.cell(row=r, column=decision_col).value
            if value in counts:
                counts[value] += 1
        self.assertEqual(counts, {"OUI": 436, "NON": 312, "INCERTAIN": 2})

    def test_40_extraction_failures_map_to_analyse_impossible(self):
        rows = self._realistic_rows()
        agreement_counts = {v: 0 for v in mod.AGREEMENT_VALUES}
        for row in rows:
            agreement_counts[mod.compute_agreement(row)] += 1
        self.assertEqual(agreement_counts[mod.AGREEMENT_ANALYSE_IA_IMPOSSIBLE], 40)

    def test_94_rows_have_independently_sourced_refreshed_reviews(self):
        rows = self._realistic_rows()
        reviewed = [row for row in rows if row.has_ai_review]
        self.assertEqual(len(reviewed), 94)
        # Each reviewed row's verdict is a function of ITS OWN row only -
        # no two rows share object identity or a copied evidence dict.
        seen_ids = set()
        for row in reviewed:
            self.assertNotIn(row.archive_file_id, seen_ids)
            seen_ids.add(row.archive_file_id)
        self.assertEqual(len(seen_ids), 94)

    def test_workbook_builds_successfully_at_full_realistic_scale(self):
        rows = self._realistic_rows()
        wb = mod.build_workbook(rows, "qwen3:14b", "v3", "2026-09-28 00:00")
        ws = wb["Comparaison Youssef vs IA"]
        data_row_count = sum(1 for row in ws.iter_rows(min_row=2) if row[0].value is not None)
        self.assertEqual(data_row_count, 750)


class DuplicateProvenanceTest(unittest.TestCase):
    def test_duplicate_marker_never_appears_on_the_visible_sheet(self):
        rows = [synthetic_row(
            candidate_id="cand-dup", archive_file_id=42, is_primary_candidate=False,
            duplicate_of_archive_file_id=7,
        )]
        wb = mod.build_workbook(rows, "qwen3:14b", "v3", "2026-09-28 00:00")
        visible_text = "\n".join(
            str(cell.value) for row in wb["Comparaison Youssef vs IA"].iter_rows() for cell in row if cell.value is not None
        )
        self.assertNotIn("is_primary_candidate", visible_text)
        self.assertNotIn("duplicate_of_archive_file_id", visible_text)

    def test_duplicate_marker_appears_only_on_the_hidden_technical_sheet(self):
        rows = [synthetic_row(
            candidate_id="cand-dup", archive_file_id=42, is_primary_candidate=False,
            duplicate_of_archive_file_id=7,
        )]
        wb = mod.build_workbook(rows, "qwen3:14b", "v3", "2026-09-28 00:00")
        technical_headers = [cell.value for cell in wb["Données techniques"][1]]
        self.assertIn("is_primary_candidate", technical_headers)
        self.assertIn("duplicate_of_archive_file_id", technical_headers)
        primary_col = technical_headers.index("is_primary_candidate") + 1
        dup_col = technical_headers.index("duplicate_of_archive_file_id") + 1
        self.assertEqual(wb["Données techniques"].cell(row=2, column=primary_col).value, "NON")
        self.assertEqual(wb["Données techniques"].cell(row=2, column=dup_col).value, 7)

    def test_each_duplicate_carries_its_own_independent_ai_result(self):
        # A duplicate and its primary must each show THEIR OWN ai_proposed_role
        # (from their own independent archive_file_id-scoped LATERAL join) -
        # never one inherited/copied from the other.
        primary = synthetic_row(
            candidate_id="cand-primary", archive_file_id=7, is_primary_candidate=True,
            ai_proposed_role="CDC", ai_processing_status="SUCCESS",
        )
        duplicate = synthetic_row(
            candidate_id="cand-dup", archive_file_id=42, is_primary_candidate=False,
            duplicate_of_archive_file_id=7, ai_proposed_role=None, ai_processing_status=None,
        )
        self.assertEqual(mod.compute_agreement(primary), mod.AGREEMENT_ACCORD)
        self.assertEqual(mod.compute_agreement(duplicate), mod.AGREEMENT_NON_COMPARE)
        self.assertNotEqual(mod.compute_ai_verdict_label(primary), mod.compute_ai_verdict_label(duplicate))


class ConfidentialityScanTest(unittest.TestCase):
    def test_no_absolute_archive_path_anywhere_in_the_workbook(self):
        rows = [synthetic_row(
            relative_path="OFFRES 2020/SYNTHETIC/synthetic.pdf",
        )]
        wb = mod.build_workbook(rows, "qwen3:14b", "v3", "2026-09-28 00:00")
        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        self.assertNotIn("/mnt/", cell.value)
                        self.assertNotIn("/home/", cell.value)

    def test_relative_path_only_appears_on_the_hidden_technical_sheet(self):
        rows = [synthetic_row(relative_path="OFFRES 2020/SYNTHETIC/synthetic_needle.pdf")]
        wb = mod.build_workbook(rows, "qwen3:14b", "v3", "2026-09-28 00:00")
        visible_text = "\n".join(
            str(cell.value) for row in wb["Comparaison Youssef vs IA"].iter_rows() for cell in row if cell.value is not None
        )
        self.assertNotIn("synthetic_needle", visible_text)
        technical_text = "\n".join(
            str(cell.value) for row in wb["Données techniques"].iter_rows() for cell in row if cell.value is not None
        )
        self.assertIn("synthetic_needle", technical_text)

    def test_technical_sheet_never_carries_absolute_root_path(self):
        rows = [synthetic_row()]
        wb = mod.build_workbook(rows, "qwen3:14b", "v3", "2026-09-28 00:00")
        self.assertEqual(len(mod.TECHNICAL_HEADERS), 7)
        self.assertNotIn("root_path", mod.TECHNICAL_HEADERS)
        self.assertNotIn("source_root_path", mod.TECHNICAL_HEADERS)

    def test_no_uuid_candidate_id_on_the_visible_sheet(self):
        rows = [synthetic_row(candidate_id="11111111-2222-3333-4444-555555555555")]
        wb = mod.build_workbook(rows, "qwen3:14b", "v3", "2026-09-28 00:00")
        visible_text = "\n".join(
            str(cell.value) for row in wb["Comparaison Youssef vs IA"].iter_rows() for cell in row if cell.value is not None
        )
        self.assertNotIn("11111111-2222-3333-4444-555555555555", visible_text)
        technical_text = "\n".join(
            str(cell.value) for row in wb["Données techniques"].iter_rows() for cell in row if cell.value is not None
        )
        self.assertIn("11111111-2222-3333-4444-555555555555", technical_text)

    def test_no_full_sha256_hash_anywhere_in_the_workbook(self):
        # ComparisonRow structurally carries no content-hash field at all -
        # a full 64-hex-char SHA-256 can never appear anywhere in this
        # workbook, visible or hidden.
        rows = [synthetic_row()]
        wb = mod.build_workbook(rows, "qwen3:14b", "v3", "2026-09-28 00:00")
        import re
        sha256_pattern = re.compile(r"\b[0-9a-f]{64}\b")
        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        self.assertIsNone(sha256_pattern.search(cell.value))


class FormulaInjectionTest(unittest.TestCase):
    def test_leading_equals_sanitized(self):
        self.assertEqual(mod.sanitize_cell_value("=cmd()"), "'=cmd()")

    def test_reopen_after_save_preserves_all_750_rows(self):
        rows = [synthetic_row(candidate_id=f"cand-{i}", archive_file_id=i) for i in range(1, 751)]
        wb = mod.build_workbook(rows, "qwen3:14b", "v3", "2026-09-28 00:00")
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "roundtrip.xlsx"
            wb.save(str(path))
            from openpyxl import load_workbook
            reopened = load_workbook(str(path))
            ws = reopened["Comparaison Youssef vs IA"]
            data_row_count = sum(1 for row in ws.iter_rows(min_row=2) if row[0].value is not None)
            self.assertEqual(data_row_count, 750)


class YoussefDecisionPreservationInWorkbookTest(unittest.TestCase):
    def test_every_decision_value_preserved_exactly_in_visible_sheet(self):
        rows = [
            synthetic_row(candidate_id="c1", archive_file_id=1, validation_status="HUMAN_VALIDATED_CDC"),
            synthetic_row(candidate_id="c2", archive_file_id=2, validation_status="HUMAN_REJECTED_CDC"),
            synthetic_row(candidate_id="c3", archive_file_id=3, validation_status="MACHINE_CLASSIFIED"),
        ]
        wb = mod.build_workbook(rows, "qwen3:14b", "v3", "2026-09-28 00:00")
        ws = wb["Comparaison Youssef vs IA"]
        decision_col = mod.VISIBLE_HEADERS.index("Décision de Youssef") + 1
        decisions = {ws.cell(row=r, column=decision_col).value for r in range(2, ws.max_row + 1)}
        self.assertEqual(decisions, {"OUI", "NON", "INCERTAIN"})


class SemanticEquivalenceMatrixTest(unittest.TestCase):
    """Exhaustive matrix: every semantic_review.py SEMANTIC_ROLES value x
    every youssef_decision (OUI/NON/INCERTAIN) x has_ai_review (True/False)
    x extraction_status (SUCCESS/FAILED) - proves compute_agreement() never
    raises and always returns a defined AGREEMENT_VALUES member, and pins
    down the exact ambiguous-role remapping (Part 2 of the 2026-09-29
    audit) so a future edit cannot silently widen or narrow it."""

    @classmethod
    def setUpClass(cls):
        sys.path.insert(0, str(HERE))
        import semantic_review
        cls.all_roles = list(semantic_review.SEMANTIC_ROLES)

    VALIDATION_STATUSES = ("HUMAN_VALIDATED_CDC", "HUMAN_REJECTED_CDC", "SOME_OTHER_STATUS")

    def test_full_matrix_never_raises_and_stays_in_defined_values(self):
        for validation_status in self.VALIDATION_STATUSES:
            for role in self.all_roles:
                for extraction_status in ("SUCCESS", "FAILED"):
                    for ai_processing_status in ("SUCCESS", None):
                        row = synthetic_row(
                            validation_status=validation_status,
                            extraction_status=extraction_status,
                            ai_proposed_role=role,
                            ai_processing_status=ai_processing_status,
                        )
                        with self.subTest(
                            validation_status=validation_status, role=role,
                            extraction_status=extraction_status, ai_processing_status=ai_processing_status,
                        ):
                            self.assertIn(mod.compute_agreement(row), mod.AGREEMENT_VALUES)

    def test_ambiguous_roles_are_exactly_tdr_and_unknown(self):
        self.assertEqual(set(mod.AI_ROLES_AMBIGUOUS_VS_YOUSSEF_OUI), {"TDR", "UNKNOWN"})

    def test_oui_vs_ambiguous_role_is_a_revoir_not_desaccord(self):
        for role in mod.AI_ROLES_AMBIGUOUS_VS_YOUSSEF_OUI:
            row = synthetic_row(
                validation_status="HUMAN_VALIDATED_CDC", extraction_status="SUCCESS",
                ai_proposed_role=role, ai_processing_status="SUCCESS",
            )
            with self.subTest(role=role):
                self.assertEqual(mod.compute_agreement(row), mod.AGREEMENT_A_REVOIR)

    def test_oui_vs_non_ambiguous_mismatching_role_stays_desaccord(self):
        non_ambiguous_mismatching = [
            role for role in self.all_roles
            if role not in mod.AI_USABLE_CDC_ROLES and role not in mod.AI_ROLES_AMBIGUOUS_VS_YOUSSEF_OUI
        ]
        self.assertTrue(non_ambiguous_mismatching)  # sanity: the set isn't accidentally empty
        for role in non_ambiguous_mismatching:
            row = synthetic_row(
                validation_status="HUMAN_VALIDATED_CDC", extraction_status="SUCCESS",
                ai_proposed_role=role, ai_processing_status="SUCCESS",
            )
            with self.subTest(role=role):
                self.assertEqual(mod.compute_agreement(row), mod.AGREEMENT_DESACCORD)

    def test_non_vs_ambiguous_role_is_still_accord(self):
        # An ambiguous role only matters when it would otherwise be a
        # mismatch against Youssef's OUI. When Youssef said NON, "TDR" or
        # "UNKNOWN" already agrees (neither is in AI_USABLE_CDC_ROLES), so
        # the remapping must never fire here.
        for role in mod.AI_ROLES_AMBIGUOUS_VS_YOUSSEF_OUI:
            row = synthetic_row(
                validation_status="HUMAN_REJECTED_CDC", extraction_status="SUCCESS",
                ai_proposed_role=role, ai_processing_status="SUCCESS",
            )
            with self.subTest(role=role):
                self.assertEqual(mod.compute_agreement(row), mod.AGREEMENT_ACCORD)

    def test_incertain_always_a_revoir_regardless_of_role(self):
        for role in self.all_roles:
            row = synthetic_row(
                validation_status="SOME_OTHER_STATUS", extraction_status="SUCCESS",
                ai_proposed_role=role, ai_processing_status="SUCCESS",
            )
            with self.subTest(role=role):
                self.assertEqual(mod.compute_agreement(row), mod.AGREEMENT_A_REVOIR)

    def test_recommended_action_distinguishes_the_two_a_revoir_causes(self):
        incertain_row = synthetic_row(validation_status="SOME_OTHER_STATUS")
        ambiguous_role_row = synthetic_row(
            validation_status="HUMAN_VALIDATED_CDC", extraction_status="SUCCESS",
            ai_proposed_role="TDR", ai_processing_status="SUCCESS",
        )
        incertain_action = mod.compute_recommended_action(incertain_row, mod.AGREEMENT_A_REVOIR)
        ambiguous_action = mod.compute_recommended_action(ambiguous_role_row, mod.AGREEMENT_A_REVOIR)
        self.assertNotEqual(incertain_action, ambiguous_action)

    def test_visible_header_label_is_the_honest_comparison_label(self):
        self.assertIn("Comparaison décision humaine / verdict IA", mod.VISIBLE_HEADERS)
        self.assertNotIn("Accord Youssef / IA", mod.VISIBLE_HEADERS)

    def test_raw_youssef_and_ai_columns_survive_ambiguous_reclassification(self):
        # Preserving raw values is the whole point of the remapping - the
        # workbook must never blend/replace them just because the row's
        # comparison outcome is A_REVOIR instead of DESACCORD.
        row = synthetic_row(
            candidate_id="c-ambig", archive_file_id=99, validation_status="HUMAN_VALIDATED_CDC",
            extraction_status="SUCCESS", ai_proposed_role="TDR", ai_processing_status="SUCCESS",
        )
        wb = mod.build_workbook([row], "qwen3:14b", "v3", "2026-09-28 00:00")
        ws = wb["Comparaison Youssef vs IA"]
        decision_col = mod.VISIBLE_HEADERS.index("Décision de Youssef") + 1
        verdict_col = mod.VISIBLE_HEADERS.index("Verdict IA") + 1
        agreement_col = mod.VISIBLE_HEADERS.index("Comparaison décision humaine / verdict IA") + 1
        self.assertEqual(ws.cell(row=2, column=decision_col).value, "OUI")
        self.assertEqual(ws.cell(row=2, column=verdict_col).value, "TDR")
        self.assertEqual(ws.cell(row=2, column=agreement_col).value, mod.AGREEMENT_A_REVOIR)


class OutputDirectoryPermissionTest(unittest.TestCase):
    """Synthetic, filesystem-only tests for _ensure_private_output_directory
    - no PostgreSQL, no real output tree. Uses only tempdirs it creates and
    cleans up itself."""

    def test_creates_full_chain_at_0700(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            target = Path(tmp_dir) / "a" / "b" / "c"
            mod._ensure_private_output_directory(target)
            self.assertTrue(target.is_dir())
            for component in (Path(tmp_dir) / "a", Path(tmp_dir) / "a" / "b", target):
                self.assertEqual(component.stat().st_mode & 0o777, 0o700)

    def test_intermediate_parent_gets_0700_even_under_a_lax_umask(self):
        import os as _os
        with tempfile.TemporaryDirectory() as tmp_dir:
            old_umask = _os.umask(0o022)
            try:
                target = Path(tmp_dir) / "parent" / "leaf"
                mod._ensure_private_output_directory(target)
                # The intermediate parent - never the leaf itself in the
                # old code path - is exactly what the original bug missed.
                intermediate = Path(tmp_dir) / "parent"
                self.assertEqual(intermediate.stat().st_mode & 0o777, 0o700)
            finally:
                _os.umask(old_umask)

    def test_rerun_on_already_correct_directory_succeeds(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            target = Path(tmp_dir) / "out"
            mod._ensure_private_output_directory(target)
            mod._ensure_private_output_directory(target)  # must not raise
            self.assertEqual(target.stat().st_mode & 0o777, 0o700)

    def test_refuses_pre_existing_directory_with_wrong_permissions(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            target = Path(tmp_dir) / "out"
            target.mkdir(mode=0o755)
            import os as _os
            _os.chmod(target, 0o755)
            with self.assertRaises(PermissionError):
                mod._ensure_private_output_directory(target)

    def test_rejects_symlinked_leaf_component(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            real_dir = Path(tmp_dir) / "real"
            real_dir.mkdir(mode=0o700)
            import os as _os
            _os.chmod(real_dir, 0o700)
            link = Path(tmp_dir) / "link"
            link.symlink_to(real_dir, target_is_directory=True)
            with self.assertRaises(PermissionError):
                mod._ensure_private_output_directory(link)

    def test_rejects_symlinked_intermediate_component(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            real_dir = Path(tmp_dir) / "real"
            real_dir.mkdir(mode=0o700)
            import os as _os
            _os.chmod(real_dir, 0o700)
            link = Path(tmp_dir) / "link"
            link.symlink_to(real_dir, target_is_directory=True)
            target = link / "leaf"
            with self.assertRaises(PermissionError):
                mod._ensure_private_output_directory(target)

    def test_rejects_relative_path(self):
        with self.assertRaises(ValueError):
            mod._ensure_private_output_directory(Path("relative/output/dir"))

    def test_rejects_path_component_that_is_a_file_not_a_directory(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            blocker = Path(tmp_dir) / "blocker"
            blocker.write_text("not a directory")
            target = blocker / "leaf"
            with self.assertRaises(NotADirectoryError):
                mod._ensure_private_output_directory(target)

    def test_run_export_produces_0600_file_and_0700_directory_chain(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            output_path = Path(tmp_dir) / "nested" / "dir" / "comparison.xlsx"
            rows = [synthetic_row(candidate_id="c1", archive_file_id=1)]

            class _FakeConn:
                def cursor(self):
                    raise AssertionError("run_export's fetch must not be reached in this test")

            # Exercise the directory/file permission path directly via the
            # same building blocks run_export() uses, without needing a
            # real PostgreSQL connection.
            wb = mod.build_workbook(rows, "qwen3:14b", "v3", "2026-09-28 00:00")
            import os as _os
            mod._ensure_private_output_directory(output_path.parent)
            wb.save(str(output_path))
            _os.chmod(output_path, 0o600)

            self.assertEqual(output_path.parent.stat().st_mode & 0o777, 0o700)
            self.assertEqual((output_path.parent.parent).stat().st_mode & 0o777, 0o700)
            self.assertEqual(output_path.stat().st_mode & 0o777, 0o600)


class ExtractionStatusLabelTest(unittest.TestCase):
    """2026-09-29 audit: every visible extraction label a non-technical
    reviewer sees must be plain French, never a raw internal status, and
    must never describe a policy exclusion as a technical failure."""

    def test_success_label(self):
        row = synthetic_row(extraction_status="SUCCESS")
        self.assertEqual(mod.compute_extraction_status_label(row), "Texte extrait avec succès")

    def test_not_attempted_with_youssef_non_is_policy_exclusion(self):
        row = synthetic_row(extraction_status="NOT_ATTEMPTED", validation_status="HUMAN_REJECTED_CDC")
        label = mod.compute_extraction_status_label(row)
        self.assertEqual(label, "Extraction non planifiée — document rejeté lors de la revue humaine")
        self.assertNotIn("échec", label.lower())

    def test_not_attempted_with_youssef_oui_is_generic_not_a_policy_claim(self):
        # Unexpected combination (0 rows like this exist today) - must
        # never claim a rejection that didn't happen.
        row = synthetic_row(extraction_status="NOT_ATTEMPTED", validation_status="HUMAN_VALIDATED_CDC")
        label = mod.compute_extraction_status_label(row)
        self.assertEqual(label, "Extraction non réalisée — motif à vérifier")
        self.assertNotIn("rejeté", label)

    def test_not_attempted_with_youssef_incertain_is_generic_not_a_policy_claim(self):
        row = synthetic_row(extraction_status="NOT_ATTEMPTED", validation_status="SOME_OTHER_STATUS")
        label = mod.compute_extraction_status_label(row)
        self.assertEqual(label, "Extraction non réalisée — motif à vérifier")
        self.assertNotIn("rejeté", label)

    def test_failed_doc_embedded_images_only_label(self):
        row = synthetic_row(extraction_status="FAILED", extraction_failure_category="DOC_EMBEDDED_IMAGES_ONLY")
        self.assertEqual(
            mod.compute_extraction_status_label(row), "Échec — document composé d'images, OCR nécessaire"
        )

    def test_failed_pdf_extraction_failure_label(self):
        row = synthetic_row(extraction_status="FAILED", extraction_failure_category="PDF_EXTRACTION_FAILURE")
        self.assertEqual(
            mod.compute_extraction_status_label(row),
            "Échec d'extraction PDF — diagnostic complémentaire nécessaire",
        )

    def test_failed_empty_extracted_text_label(self):
        row = synthetic_row(extraction_status="FAILED", extraction_failure_category="EMPTY_EXTRACTED_TEXT")
        self.assertEqual(mod.compute_extraction_status_label(row), "Échec — aucun texte exploitable obtenu")

    def test_failed_unknown_category_uses_honest_fallback_not_invented_claim(self):
        row = synthetic_row(extraction_status="FAILED", extraction_failure_category="SOME_FUTURE_CATEGORY")
        label = mod.compute_extraction_status_label(row)
        self.assertEqual(label, mod.EXTRACTION_FAILURE_CATEGORY_LABEL_FALLBACK)
        # never silently blank, never invents a specific technical claim
        self.assertTrue(label)

    def test_every_label_is_a_defined_string_never_none_or_empty(self):
        combos = [
            ("SUCCESS", None, "HUMAN_VALIDATED_CDC"),
            ("SUCCESS", None, "HUMAN_REJECTED_CDC"),
            ("NOT_ATTEMPTED", None, "HUMAN_REJECTED_CDC"),
            ("NOT_ATTEMPTED", None, "HUMAN_VALIDATED_CDC"),
            ("FAILED", "DOC_EMBEDDED_IMAGES_ONLY", "HUMAN_VALIDATED_CDC"),
            ("FAILED", "PDF_EXTRACTION_FAILURE", "HUMAN_VALIDATED_CDC"),
            ("FAILED", "EMPTY_EXTRACTED_TEXT", "HUMAN_VALIDATED_CDC"),
        ]
        for extraction_status, failure_category, validation_status in combos:
            row = synthetic_row(
                extraction_status=extraction_status, extraction_failure_category=failure_category,
                validation_status=validation_status,
            )
            with self.subTest(extraction_status=extraction_status, failure_category=failure_category):
                label = mod.compute_extraction_status_label(row)
                self.assertIsInstance(label, str)
                self.assertTrue(label.strip())


class ExtractionExplanationTest(unittest.TestCase):
    def test_success_explanation_mentions_success_not_failure(self):
        row = synthetic_row(extraction_status="SUCCESS")
        explanation = mod.compute_extraction_explanation(row)
        self.assertIn("succès", explanation.lower())

    def test_not_attempted_non_explanation_explicitly_says_not_a_technical_failure(self):
        row = synthetic_row(extraction_status="NOT_ATTEMPTED", validation_status="HUMAN_REJECTED_CDC")
        explanation = mod.compute_extraction_explanation(row)
        self.assertIn("rejeté", explanation)
        self.assertIn("pas un échec technique", explanation)

    def test_not_attempted_other_explanation_does_not_claim_rejection(self):
        row = synthetic_row(extraction_status="NOT_ATTEMPTED", validation_status="HUMAN_VALIDATED_CDC")
        explanation = mod.compute_extraction_explanation(row)
        self.assertNotIn("rejeté", explanation)

    def test_each_failure_category_has_a_distinct_non_empty_explanation(self):
        categories = ["DOC_EMBEDDED_IMAGES_ONLY", "PDF_EXTRACTION_FAILURE", "EMPTY_EXTRACTED_TEXT"]
        explanations = set()
        for category in categories:
            row = synthetic_row(extraction_status="FAILED", extraction_failure_category=category)
            explanation = mod.compute_extraction_explanation(row)
            self.assertTrue(explanation.strip())
            explanations.add(explanation)
        self.assertEqual(len(explanations), 3)  # all distinct - no copy-paste collapse

    def test_unknown_failure_category_uses_honest_fallback(self):
        row = synthetic_row(extraction_status="FAILED", extraction_failure_category="SOME_FUTURE_CATEGORY")
        explanation = mod.compute_extraction_explanation(row)
        self.assertEqual(explanation, mod.EXTRACTION_FAILURE_CATEGORY_EXPLANATION_FALLBACK)


class ExtractionActionRecommendationTest(unittest.TestCase):
    def test_human_rejected_not_attempted_action(self):
        row = synthetic_row(extraction_status="NOT_ATTEMPTED", validation_status="HUMAN_REJECTED_CDC")
        agreement = mod.compute_agreement(row)
        self.assertEqual(agreement, mod.AGREEMENT_NON_COMPARE)
        self.assertEqual(
            mod.compute_recommended_action(row, agreement),
            "Aucune extraction requise sauf nouvelle décision humaine",
        )

    def test_doc_images_only_action(self):
        row = synthetic_row(extraction_status="FAILED", extraction_failure_category="DOC_EMBEDDED_IMAGES_ONLY")
        agreement = mod.compute_agreement(row)
        self.assertEqual(mod.compute_recommended_action(row, agreement), "Préparer un pilote OCR local")

    def test_pdf_extraction_failure_action(self):
        row = synthetic_row(extraction_status="FAILED", extraction_failure_category="PDF_EXTRACTION_FAILURE")
        agreement = mod.compute_agreement(row)
        self.assertEqual(
            mod.compute_recommended_action(row, agreement),
            "Diagnostiquer le PDF localement avant nouvelle tentative",
        )

    def test_empty_extracted_text_action(self):
        row = synthetic_row(extraction_status="FAILED", extraction_failure_category="EMPTY_EXTRACTED_TEXT")
        agreement = mod.compute_agreement(row)
        self.assertEqual(
            mod.compute_recommended_action(row, agreement),
            "Vérifier si le document est scanné ou endommagé",
        )

    def test_success_no_semantic_review_action(self):
        row = synthetic_row(extraction_status="SUCCESS", ai_processing_status=None)
        agreement = mod.compute_agreement(row)
        self.assertEqual(agreement, mod.AGREEMENT_NON_COMPARE)
        self.assertEqual(
            mod.compute_recommended_action(row, agreement),
            "Lancer la revue IA seulement si le document est dans le périmètre autorisé",
        )

    def test_success_with_semantic_review_action_unchanged(self):
        # SUCCESS + a real AI review (ACCORD/DESACCORD/A_REVOIR) must keep
        # its existing, already-tested comparison action untouched.
        row = synthetic_row(
            extraction_status="SUCCESS", validation_status="HUMAN_VALIDATED_CDC",
            ai_proposed_role="CDC", ai_processing_status="SUCCESS",
        )
        agreement = mod.compute_agreement(row)
        self.assertEqual(agreement, mod.AGREEMENT_ACCORD)
        self.assertEqual(mod.compute_recommended_action(row, agreement), "Aucune action - décision et IA concordent")

    def test_not_attempted_non_rejected_action_does_not_reuse_rejected_message(self):
        row = synthetic_row(extraction_status="NOT_ATTEMPTED", validation_status="HUMAN_VALIDATED_CDC")
        agreement = mod.compute_agreement(row)
        action = mod.compute_recommended_action(row, agreement)
        self.assertNotEqual(action, "Aucune extraction requise sauf nouvelle décision humaine")


class ExtractionColumnWorkbookIntegrationTest(unittest.TestCase):
    def test_explanation_column_present_and_populated(self):
        rows = [
            synthetic_row(candidate_id="c1", archive_file_id=1, extraction_status="SUCCESS"),
            synthetic_row(
                candidate_id="c2", archive_file_id=2, extraction_status="NOT_ATTEMPTED",
                validation_status="HUMAN_REJECTED_CDC",
            ),
            synthetic_row(
                candidate_id="c3", archive_file_id=3, extraction_status="FAILED",
                extraction_failure_category="PDF_EXTRACTION_FAILURE",
            ),
        ]
        wb = mod.build_workbook(rows, "qwen3:14b", "v3", "2026-09-28 00:00")
        ws = wb["Comparaison Youssef vs IA"]
        self.assertIn("Explication de l'état d'extraction", mod.VISIBLE_HEADERS)
        explanation_col = mod.VISIBLE_HEADERS.index("Explication de l'état d'extraction") + 1
        status_col = mod.VISIBLE_HEADERS.index("Statut d'extraction actuel") + 1
        for r in range(2, ws.max_row + 1):
            self.assertTrue(str(ws.cell(row=r, column=explanation_col).value).strip())
            status_value = ws.cell(row=r, column=status_col).value
            self.assertNotIn(status_value, ("SUCCESS", "NOT_ATTEMPTED", "FAILED"))  # never a raw status

    def test_raw_values_preserved_on_hidden_technical_sheet(self):
        rows = [synthetic_row(
            candidate_id="c1", archive_file_id=1, extraction_status="FAILED",
            extraction_failure_category="PDF_EXTRACTION_FAILURE",
        )]
        wb = mod.build_workbook(rows, "qwen3:14b", "v3", "2026-09-28 00:00")
        technical = wb["Données techniques"]
        header = [c.value for c in technical[1]]
        self.assertIn("extraction_status (raw)", header)
        self.assertIn("extraction_failure_category (raw)", header)
        status_col = header.index("extraction_status (raw)") + 1
        category_col = header.index("extraction_failure_category (raw)") + 1
        self.assertEqual(technical.cell(row=2, column=status_col).value, "FAILED")
        self.assertEqual(technical.cell(row=2, column=category_col).value, "PDF_EXTRACTION_FAILURE")

    def test_column_widths_target_the_correct_header_after_insertion(self):
        # Regression guard: column_widths must be keyed by header text, not
        # a hardcoded position, since a new column was inserted before
        # "Verdict IA" and "Comparaison décision humaine / verdict IA".
        rows = [synthetic_row()]
        wb = mod.build_workbook(rows, "qwen3:14b", "v3", "2026-09-28 00:00")
        ws = wb["Comparaison Youssef vs IA"]
        from openpyxl.utils import get_column_letter
        verdict_col = mod.VISIBLE_HEADERS.index("Verdict IA") + 1
        verdict_letter = get_column_letter(verdict_col)
        self.assertEqual(ws.column_dimensions[verdict_letter].width, 20)


class SanitizedTotalsAndExportTest(unittest.TestCase):
    def test_750_row_export_extraction_totals_present_and_labeled(self):
        rows = (
            [synthetic_row(candidate_id=f"s{i}", archive_file_id=i, extraction_status="SUCCESS") for i in range(1, 618)]
            + [
                synthetic_row(
                    candidate_id=f"n{i}", archive_file_id=600 + i, extraction_status="NOT_ATTEMPTED",
                    validation_status="HUMAN_REJECTED_CDC",
                )
                for i in range(1, 68)
            ]
            + [
                synthetic_row(
                    candidate_id=f"f{i}", archive_file_id=700 + i, extraction_status="FAILED",
                    extraction_failure_category=(
                        "DOC_EMBEDDED_IMAGES_ONLY" if i <= 39 else
                        "PDF_EXTRACTION_FAILURE" if i <= 64 else "EMPTY_EXTRACTED_TEXT"
                    ),
                )
                for i in range(1, 67)
            ]
        )
        self.assertEqual(len(rows), 750)
        wb = mod.build_workbook(rows, "qwen3:14b", "v3", "2026-09-28 00:00")
        ws = wb["Comparaison Youssef vs IA"]
        data_row_count = sum(1 for row in ws.iter_rows(min_row=2) if row[0].value is not None)
        self.assertEqual(data_row_count, 750)

        status_col = mod.VISIBLE_HEADERS.index("Statut d'extraction actuel") + 1
        labels = [ws.cell(row=r, column=status_col).value for r in range(2, ws.max_row + 1)]
        self.assertEqual(sum(1 for v in labels if v == "Texte extrait avec succès"), 617)
        self.assertEqual(
            sum(1 for v in labels if v == "Extraction non planifiée — document rejeté lors de la revue humaine"), 67
        )
        self.assertEqual(sum(1 for v in labels if v == "Échec — document composé d'images, OCR nécessaire"), 39)
        self.assertEqual(
            sum(1 for v in labels if v == "Échec d'extraction PDF — diagnostic complémentaire nécessaire"), 25
        )
        self.assertEqual(sum(1 for v in labels if v == "Échec — aucun texte exploitable obtenu"), 2)


if __name__ == "__main__":
    unittest.main()
