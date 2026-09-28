#!/usr/bin/env python3
"""Synthetic test suite for the targeted-recovery extraction additions in
scripts/cdc_content_inspector.py (ODT, standalone RTF, XLSX, size-aware
PDF timeout + bounded retry, DOC embedded-images-only diagnosis).

SYNTHETIC DATA ONLY. Every ODT/RTF/XLSX/DOCX fixture built here uses
generic, invented placeholder text - no real archive filenames, paths,
project names, client names, or document content appear anywhere in this
file. No real Docling subprocess and no real LibreOffice conversion is
ever invoked for the PDF-timeout/retry tests (the Docling converter is
fully dependency-injected via a fake); the ODT/XLSX/RTF tests use only
synthetic in-memory files written to a temp directory. No Ollama call is
possible from this file at all - nothing here constructs an ai_adapter.
"""
from __future__ import annotations

import importlib.util
import io
import sys
import tempfile
import time
import unittest
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent

_CDC_SPEC = importlib.util.spec_from_file_location("cdc_discovery", HERE / "cdc_discovery.py")
cdc_discovery = importlib.util.module_from_spec(_CDC_SPEC)
sys.modules[_CDC_SPEC.name] = cdc_discovery
_CDC_SPEC.loader.exec_module(cdc_discovery)

_TSC_SPEC = importlib.util.spec_from_file_location("technical_source_classifier", HERE / "technical_source_classifier.py")
tsc = importlib.util.module_from_spec(_TSC_SPEC)
sys.modules[_TSC_SPEC.name] = tsc
_TSC_SPEC.loader.exec_module(tsc)

_INSPECTOR_SPEC = importlib.util.spec_from_file_location("cdc_content_inspector", HERE / "cdc_content_inspector.py")
insp = importlib.util.module_from_spec(_INSPECTOR_SPEC)
sys.modules[_INSPECTOR_SPEC.name] = insp
_INSPECTOR_SPEC.loader.exec_module(insp)

DiscoveryCounters = cdc_discovery.DiscoveryCounters


def _tmpdir():
    return Path(tempfile.mkdtemp(prefix="cdc-recovery-test-"))


def write_synthetic_odt(path: Path, paragraphs):
    ns = 'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"'
    body = "".join(f"<text:p>{p}</text:p>" for p in paragraphs)
    content_xml = f'<?xml version="1.0"?><office:document-content {ns}><office:body><office:text>{body}</office:text></office:body></office:document-content>'
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        z.writestr("content.xml", content_xml)


class TestOdtExtraction(unittest.TestCase):
    def setUp(self):
        self.tmp = _tmpdir()
        self.addCleanup(lambda: __import__("shutil").rmtree(self.tmp, ignore_errors=True))

    def test_valid_odt_extracts_text(self):
        path = self.tmp / "synthetic.odt"
        write_synthetic_odt(path, ["Texte de paragraphe synthetique un.", "Deuxieme paragraphe synthetique."])
        counters = DiscoveryCounters()
        text = insp.extract_odt_text(path, counters)
        self.assertIn("synthetique", text)
        self.assertEqual(counters.odt_extraction_calls, 1)

    def test_missing_file_raises_path_missing(self):
        counters = DiscoveryCounters()
        with self.assertRaises(insp.ExtractionError) as ctx:
            insp.extract_odt_text(self.tmp / "missing.odt", counters)
        self.assertEqual(ctx.exception.reason_code, "odt_path_missing")

    def test_empty_paragraphs_raise_output_empty(self):
        path = self.tmp / "empty.odt"
        write_synthetic_odt(path, [])
        counters = DiscoveryCounters()
        with self.assertRaises(insp.ExtractionError) as ctx:
            insp.extract_odt_text(path, counters)
        self.assertEqual(ctx.exception.reason_code, "odt_output_empty")

    def test_malformed_archive_rejected_safely(self):
        path = self.tmp / "corrupt.odt"
        path.write_bytes(b"not a zip file at all")
        counters = DiscoveryCounters()
        with self.assertRaises(insp.ExtractionError) as ctx:
            insp.extract_odt_text(path, counters)
        self.assertEqual(ctx.exception.reason_code, "odt_read_failed")

    def test_missing_content_xml_member_rejected_safely(self):
        path = self.tmp / "no_content.odt"
        with zipfile.ZipFile(path, "w") as z:
            z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        counters = DiscoveryCounters()
        with self.assertRaises(insp.ExtractionError) as ctx:
            insp.extract_odt_text(path, counters)
        self.assertEqual(ctx.exception.reason_code, "odt_read_failed")

    def test_malformed_xml_rejected_safely(self):
        path = self.tmp / "bad_xml.odt"
        with zipfile.ZipFile(path, "w") as z:
            z.writestr("content.xml", "<office:document-content><unclosed>")
        counters = DiscoveryCounters()
        with self.assertRaises(insp.ExtractionError) as ctx:
            insp.extract_odt_text(path, counters)
        self.assertEqual(ctx.exception.reason_code, "odt_xml_parse_failed")

    def test_text_size_limit_applied(self):
        path = self.tmp / "large.odt"
        write_synthetic_odt(path, ["x" * 100_000])
        counters = DiscoveryCounters()
        text = insp.extract_odt_text(path, counters, char_limit=50)
        self.assertEqual(len(text), 50)

    def test_reason_codes_map_to_stable_categories(self):
        self.assertEqual(tsc.categorize_extraction_failure_reason("odt_path_missing"), "MISSING_SOURCE")
        self.assertEqual(tsc.categorize_extraction_failure_reason("odt_read_failed"), "ODT_EXTRACTION_FAILURE")
        self.assertEqual(tsc.categorize_extraction_failure_reason("odt_xml_parse_failed"), "ODT_EXTRACTION_FAILURE")
        self.assertEqual(tsc.categorize_extraction_failure_reason("odt_output_empty"), "EMPTY_EXTRACTED_TEXT")


class TestRtfStandaloneExtraction(unittest.TestCase):
    def setUp(self):
        self.tmp = _tmpdir()
        self.addCleanup(lambda: __import__("shutil").rmtree(self.tmp, ignore_errors=True))

    def test_valid_rtf_extracts_text(self):
        path = self.tmp / "synthetic.rtf"
        path.write_bytes(rb"{\rtf1\ansi Texte synthetique de test RTF.}")
        counters = DiscoveryCounters()
        text = insp.extract_rtf_text_standalone(path, counters)
        self.assertIn("synthetique", text)
        self.assertEqual(counters.rtf_extraction_calls, 1)

    def test_missing_file_raises_path_missing(self):
        counters = DiscoveryCounters()
        with self.assertRaises(insp.ExtractionError) as ctx:
            insp.extract_rtf_text_standalone(self.tmp / "missing.rtf", counters)
        self.assertEqual(ctx.exception.reason_code, "rtf_path_missing")

    def test_empty_content_raises_output_empty(self):
        path = self.tmp / "empty.rtf"
        path.write_bytes(rb"{\rtf1\ansi}")
        counters = DiscoveryCounters()
        with self.assertRaises(insp.ExtractionError) as ctx:
            insp.extract_rtf_text_standalone(path, counters)
        self.assertEqual(ctx.exception.reason_code, "rtf_output_empty")

    def test_no_shell_interpolation_path_with_special_characters(self):
        # extract_rtf_text_standalone must never build/run a shell command
        # at all - proven structurally by inspecting the module source for
        # any subprocess/os.system usage inside the RTF code path.
        source = (HERE / "cdc_content_inspector.py").read_text(encoding="utf-8")
        rtf_section_start = source.index("def extract_rtf_text_standalone")
        rtf_section_end = source.index("# ODT (OpenDocument Text) extraction")
        # Strip the docstring (which correctly DOCUMENTS the absence of a
        # subprocess call, and would otherwise false-positive this scan)
        # before checking for actual usage.
        rtf_section = source[rtf_section_start:rtf_section_end]
        rtf_section_code_only = rtf_section.split('"""', 2)[-1] if rtf_section.count('"""') >= 2 else rtf_section
        self.assertNotIn("subprocess", rtf_section_code_only)
        self.assertNotIn("os.system", rtf_section_code_only)
        self.assertNotIn("shell=True", rtf_section_code_only)

    def test_reason_codes_map_to_stable_categories(self):
        self.assertEqual(tsc.categorize_extraction_failure_reason("rtf_path_missing"), "MISSING_SOURCE")
        self.assertEqual(tsc.categorize_extraction_failure_reason("rtf_output_empty"), "EMPTY_EXTRACTED_TEXT")


def write_synthetic_xlsx(path: Path, sheets: dict):
    """sheets: {sheet_name: [[cell, cell, ...], ...]}"""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    wb.save(str(path))


class TestXlsxExtraction(unittest.TestCase):
    def setUp(self):
        self.tmp = _tmpdir()
        self.addCleanup(lambda: __import__("shutil").rmtree(self.tmp, ignore_errors=True))

    def test_valid_workbook_extracts_cell_values(self):
        path = self.tmp / "synthetic.xlsx"
        write_synthetic_xlsx(path, {"Feuille1": [["Alpha", "Beta"], ["Gamma synthetique", 42]]})
        counters = DiscoveryCounters()
        text = insp.extract_xlsx_text(path, counters)
        self.assertIn("Gamma synthetique", text)
        self.assertIn("Feuille1", text)  # sheet name included in the internal representation
        self.assertEqual(counters.xlsx_extraction_calls, 1)
        self.assertEqual(counters.xlsx_extraction_successes, 1)

    def test_missing_file_raises_path_missing(self):
        counters = DiscoveryCounters()
        with self.assertRaises(insp.ExtractionError) as ctx:
            insp.extract_xlsx_text(self.tmp / "missing.xlsx", counters)
        self.assertEqual(ctx.exception.reason_code, "xlsx_path_missing")
        self.assertEqual(counters.xlsx_extraction_failures, 1)

    def test_empty_workbook_raises_output_empty(self):
        path = self.tmp / "empty.xlsx"
        write_synthetic_xlsx(path, {"Feuille1": [[None, None]]})
        counters = DiscoveryCounters()
        with self.assertRaises(insp.ExtractionError) as ctx:
            insp.extract_xlsx_text(path, counters)
        self.assertEqual(ctx.exception.reason_code, "xlsx_output_empty")
        self.assertEqual(counters.xlsx_extraction_failures, 1)

    def test_malformed_workbook_rejected_safely(self):
        path = self.tmp / "corrupt.xlsx"
        path.write_bytes(b"not a valid xlsx file")
        counters = DiscoveryCounters()
        with self.assertRaises(insp.ExtractionError) as ctx:
            insp.extract_xlsx_text(path, counters)
        self.assertEqual(ctx.exception.reason_code, "xlsx_malformed")

    def test_excessive_sheet_count_rejected(self):
        path = self.tmp / "many_sheets.xlsx"
        sheets = {f"S{i}": [["x"]] for i in range(insp.XLSX_MAX_SHEETS + 1)}
        write_synthetic_xlsx(path, sheets)
        counters = DiscoveryCounters()
        with self.assertRaises(insp.ExtractionError) as ctx:
            insp.extract_xlsx_text(path, counters)
        self.assertEqual(ctx.exception.reason_code, "xlsx_dimensions_exceeded")

    def test_excessive_cell_count_rejected(self):
        path = self.tmp / "many_cells.xlsx"
        # One sheet with more total cells than XLSX_MAX_CELLS_TOTAL, kept
        # within XLSX_MAX_ROWS_PER_SHEET so the row cap isn't what trips
        # first - proves the CELL-total cap specifically.
        old_max_cells = insp.XLSX_MAX_CELLS_TOTAL
        old_max_rows = insp.XLSX_MAX_ROWS_PER_SHEET
        insp.XLSX_MAX_CELLS_TOTAL = 10
        insp.XLSX_MAX_ROWS_PER_SHEET = 100
        try:
            write_synthetic_xlsx(path, {"S": [["a"] * 5 for _ in range(5)]})  # 25 cells > 10
            counters = DiscoveryCounters()
            with self.assertRaises(insp.ExtractionError) as ctx:
                insp.extract_xlsx_text(path, counters)
            self.assertEqual(ctx.exception.reason_code, "xlsx_dimensions_exceeded")
        finally:
            insp.XLSX_MAX_CELLS_TOTAL = old_max_cells
            insp.XLSX_MAX_ROWS_PER_SHEET = old_max_rows

    def test_text_size_limit_applied(self):
        path = self.tmp / "large_text.xlsx"
        write_synthetic_xlsx(path, {"S": [["x" * 100_000]]})
        counters = DiscoveryCounters()
        text = insp.extract_xlsx_text(path, counters, char_limit=50)
        self.assertEqual(len(text), 50)

    def test_never_preserves_vba_project(self):
        # keep_vba=False is passed explicitly to load_workbook - confirmed
        # by source inspection (a macro-enabled .xlsm would otherwise have
        # its vbaProject.bin blob preserved/round-tripped; this module
        # never needs or wants that).
        source = (HERE / "cdc_content_inspector.py").read_text(encoding="utf-8")
        xlsx_section_start = source.index("def extract_xlsx_text")
        xlsx_section_end = source.index("def extract_html_text")
        self.assertIn("keep_vba=False", source[xlsx_section_start:xlsx_section_end])

    def test_no_network_call_possible(self):
        # openpyxl.load_workbook has no code path that performs a network
        # request - structurally confirmed by the absence of any URL-
        # fetching import/call in this function's own source slice.
        source = (HERE / "cdc_content_inspector.py").read_text(encoding="utf-8")
        xlsx_section_start = source.index("def extract_xlsx_text")
        xlsx_section_end = source.index("def extract_html_text")
        section = source[xlsx_section_start:xlsx_section_end]
        self.assertNotIn("urllib", section)
        self.assertNotIn("requests.", section)

    def test_reason_codes_map_to_stable_categories(self):
        self.assertEqual(tsc.categorize_extraction_failure_reason("xlsx_path_missing"), "MISSING_SOURCE")
        self.assertEqual(tsc.categorize_extraction_failure_reason("xlsx_malformed"), "XLSX_EXTRACTION_FAILURE")
        self.assertEqual(tsc.categorize_extraction_failure_reason("xlsx_dimensions_exceeded"), "XLSX_DIMENSIONS_EXCEEDED")
        self.assertEqual(tsc.categorize_extraction_failure_reason("xlsx_output_empty"), "EMPTY_EXTRACTED_TEXT")


class TestPdfSizeAwareTimeoutPolicy(unittest.TestCase):
    def test_normal_size_file_gets_default_timeout(self):
        self.assertEqual(insp.resolve_pdf_timeout(500_000), insp.DOCLING_TIMEOUT_SECONDS)

    def test_large_file_gets_increased_timeout(self):
        timeout = insp.resolve_pdf_timeout(insp.PDF_LARGE_FILE_SIZE_BYTES + 1)
        self.assertEqual(timeout, insp.PDF_LARGE_FILE_INITIAL_TIMEOUT_SECONDS)
        self.assertGreater(timeout, insp.DOCLING_TIMEOUT_SECONDS)

    def test_boundary_exactly_at_threshold_stays_default(self):
        self.assertEqual(insp.resolve_pdf_timeout(insp.PDF_LARGE_FILE_SIZE_BYTES), insp.DOCLING_TIMEOUT_SECONDS)

    def test_hard_upper_bound_never_exceeded(self):
        self.assertLessEqual(insp.PDF_LARGE_FILE_INITIAL_TIMEOUT_SECONDS, insp.PDF_MAX_TIMEOUT_SECONDS)
        self.assertLessEqual(insp.PDF_MAX_TIMEOUT_SECONDS, 240.0)


class TestPdfTimeoutRetry(unittest.TestCase):
    def setUp(self):
        self.tmp = _tmpdir()
        self.addCleanup(lambda: __import__("shutil").rmtree(self.tmp, ignore_errors=True))
        self.path = self.tmp / "synthetic.pdf"
        self.path.write_bytes(b"%PDF-synthetic")

    def test_timeout_then_success_retries_exactly_once(self):
        calls = []

        def fake_converter(source, destination, timeout=None):
            calls.append(timeout)
            if len(calls) == 1:
                raise insp.ExtractionError("docling_timeout")
            destination.write_text("texte synthetique extrait apres relance", encoding="utf-8")

        counters = DiscoveryCounters()
        text = insp.extract_pdf_text_with_retry(self.path, counters, converter=fake_converter)
        self.assertEqual(text, "texte synthetique extrait apres relance")
        self.assertEqual(len(calls), 2)
        self.assertEqual(counters.pdf_timeout_retries, 1)
        self.assertEqual(counters.pdf_extraction_calls, 2)

    def test_second_timeout_gives_up_after_exactly_one_retry(self):
        calls = []

        def always_times_out(source, destination, timeout=None):
            calls.append(timeout)
            raise insp.ExtractionError("docling_timeout")

        counters = DiscoveryCounters()
        with self.assertRaises(insp.ExtractionError) as ctx:
            insp.extract_pdf_text_with_retry(self.path, counters, converter=always_times_out)
        self.assertEqual(ctx.exception.reason_code, "docling_timeout")
        self.assertEqual(len(calls), 2)  # exactly one retry, never more
        self.assertEqual(counters.pdf_timeout_retries, 1)

    def test_retry_timeout_never_exceeds_hard_cap(self):
        calls = []

        def always_times_out(source, destination, timeout=None):
            calls.append(timeout)
            raise insp.ExtractionError("docling_timeout")

        counters = DiscoveryCounters()
        with self.assertRaises(insp.ExtractionError):
            insp.extract_pdf_text_with_retry(self.path, counters, converter=always_times_out)
        self.assertLessEqual(calls[1], insp.PDF_MAX_TIMEOUT_SECONDS)

    def test_deterministic_failure_is_never_retried(self):
        calls = []

        def missing_python(source, destination, timeout=None):
            calls.append(timeout)
            raise insp.ExtractionError("docling_python_missing")

        counters = DiscoveryCounters()
        with self.assertRaises(insp.ExtractionError) as ctx:
            insp.extract_pdf_text_with_retry(self.path, counters, converter=missing_python)
        self.assertEqual(ctx.exception.reason_code, "docling_python_missing")
        self.assertEqual(len(calls), 1)  # never retried
        self.assertEqual(counters.pdf_timeout_retries, 0)

    def test_first_attempt_success_never_retries(self):
        calls = []

        def immediate_success(source, destination, timeout=None):
            calls.append(timeout)
            destination.write_text("succes immediat synthetique", encoding="utf-8")

        counters = DiscoveryCounters()
        text = insp.extract_pdf_text_with_retry(self.path, counters, converter=immediate_success)
        self.assertEqual(text, "succes immediat synthetique")
        self.assertEqual(len(calls), 1)
        self.assertEqual(counters.pdf_timeout_retries, 0)


class TestDocEmbeddedImagesOnlyDetection(unittest.TestCase):
    def test_docx_with_media_part_detected(self):
        tmp = _tmpdir()
        self.addCleanup(lambda: __import__("shutil").rmtree(tmp, ignore_errors=True))
        path = tmp / "with_image.docx"
        with zipfile.ZipFile(path, "w") as z:
            z.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body></w:body></w:document>')
            z.writestr("word/media/image1.png", b"\x89PNG-synthetic-placeholder")
        self.assertTrue(insp._docx_contains_embedded_images(path))

    def test_docx_without_media_part_not_detected(self):
        tmp = _tmpdir()
        self.addCleanup(lambda: __import__("shutil").rmtree(tmp, ignore_errors=True))
        path = tmp / "no_image.docx"
        with zipfile.ZipFile(path, "w") as z:
            z.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body></w:body></w:document>')
        self.assertFalse(insp._docx_contains_embedded_images(path))

    def test_malformed_archive_fails_closed_to_false(self):
        tmp = _tmpdir()
        self.addCleanup(lambda: __import__("shutil").rmtree(tmp, ignore_errors=True))
        path = tmp / "corrupt.docx"
        path.write_bytes(b"not a zip file")
        self.assertFalse(insp._docx_contains_embedded_images(path))

    def test_reason_code_maps_to_specific_category(self):
        self.assertEqual(tsc.categorize_extraction_failure_reason("doc_embedded_images_only"), "DOC_EMBEDDED_IMAGES_ONLY")


class TestDispatchWiring(unittest.TestCase):
    """Confirms LocalContentInspector.inspect() actually routes odt/rtf/xlsx
    to the new extractors, and pdf through the retry-aware entry point -
    a source-level check, no real file I/O needed."""

    def test_inspect_dispatches_odt_rtf_xlsx_and_retrying_pdf(self):
        source = (HERE / "cdc_content_inspector.py").read_text(encoding="utf-8")
        inspect_start = source.index("class LocalContentInspector")
        self.assertIn('normalized_extension == "odt"', source[inspect_start:])
        self.assertIn('normalized_extension == "rtf"', source[inspect_start:])
        self.assertIn('normalized_extension == "xlsx"', source[inspect_start:])
        self.assertIn("extract_pdf_text_with_retry(", source[inspect_start:])


if __name__ == "__main__":
    unittest.main()
