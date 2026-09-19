#!/usr/bin/env python3
"""Synthetic test suite for cdc_review_import_preview.py.

SYNTHETIC DATA ONLY. Every workbook built here uses invented, generic
placeholder values (SYN-ID-*, ALT-ID-*, "Col3"/"Col4"/"Col5" filler header
text, "OUI"/"NON"/"INCERTAIN" decision labels, invented years). No real
archive filenames, paths, project names, identifiers, or document content
is used anywhere in this file. No PostgreSQL connection, no DATABASE_URL
read, no Ollama/Docling/n8n/Qdrant call is ever made or possible from this
test file or the module it tests.
"""
from __future__ import annotations

import importlib.util
import io
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from unittest.mock import patch

import openpyxl

HERE = Path(__file__).resolve().parent
MODULE_PATH = HERE / "cdc_review_import_preview.py"
SPEC = importlib.util.spec_from_file_location("cdc_review_import_preview", MODULE_PATH)
preview = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = preview
SPEC.loader.exec_module(preview)


# ---------------------------------------------------------------------
# Synthetic workbook builder - invented generic data only
# ---------------------------------------------------------------------

DEFAULT_MAIN_HEADERS = ["Priorité", "Année", "Col3", "Col4", "Col5", "CDC utilisable ?", "Col7"]
DEFAULT_TECHNICAL_HEADERS = [f"T{i}" for i in range(1, 25)]


def build_synthetic_workbook(
    path: Path,
    rows,  # list of dicts: {"decision": str, "annee": value, "priorite": str}
    *,
    sheet_count=5,
    main_headers=None,
    technical_headers=None,
    technical_row_count=None,  # defaults to len(rows); pass a different value to break alignment
    identifier_overrides=None,  # dict[row_index] -> (primary_id, secondary_id) override
    duplicate_primary_id_at=None,  # row index to force-duplicate with row 0's id
    blank_primary_id_at=None,  # row index to force-blank
    omit_banner=False,
):
    main_headers = main_headers if main_headers is not None else list(DEFAULT_MAIN_HEADERS)
    technical_headers = technical_headers if technical_headers is not None else list(DEFAULT_TECHNICAL_HEADERS)
    n = len(rows)
    technical_row_count = n if technical_row_count is None else technical_row_count

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Candidats"
    if not omit_banner:
        ws1.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(main_headers))
        ws1.cell(row=1, column=1, value="Synthetic Banner Title")
        header_row = 2
        data_start = 3
    else:
        header_row = 1
        data_start = 2
    for col, header in enumerate(main_headers, start=1):
        ws1.cell(row=header_row, column=col, value=header)
    for i, row in enumerate(rows):
        r = data_start + i
        ws1.cell(row=r, column=1, value=row.get("priorite", "SYN-PRIORITY-A"))
        ws1.cell(row=r, column=2, value=row.get("annee", 2021))
        ws1.cell(row=r, column=3, value="synthetic-3")
        ws1.cell(row=r, column=4, value="synthetic-4")
        ws1.cell(row=r, column=5, value="synthetic-5")
        ws1.cell(row=r, column=6, value=row.get("decision", "OUI"))
        ws1.cell(row=r, column=7, value=None)

    ws2 = wb.create_sheet("Listes")
    ws2.sheet_state = "hidden"
    ws2.cell(row=1, column=1, value="Choix")
    for i, val in enumerate(["OUI", "NON", "INCERTAIN"], start=2):
        ws2.cell(row=i, column=1, value=val)

    ws3 = wb.create_sheet("Technique")
    ws3.sheet_state = "hidden"
    if not omit_banner:
        ws3.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(technical_headers))
        ws3.cell(row=1, column=1, value="Synthetic Banner Title")
        t_header_row = 2
        t_data_start = 3
    else:
        t_header_row = 1
        t_data_start = 2
    for col, header in enumerate(technical_headers, start=1):
        ws3.cell(row=t_header_row, column=col, value=header)
    for i in range(technical_row_count):
        r = t_data_start + i
        primary_id = f"SYN-ID-{i:04d}"
        secondary_id = f"ALT-ID-{i:04d}"
        if identifier_overrides and i in identifier_overrides:
            primary_id, secondary_id = identifier_overrides[i]
        if duplicate_primary_id_at is not None and i == duplicate_primary_id_at:
            primary_id = f"SYN-ID-{0:04d}"
        if blank_primary_id_at is not None and i == blank_primary_id_at:
            primary_id = None
        ws3.cell(row=r, column=1, value=primary_id)
        ws3.cell(row=r, column=2, value=secondary_id)
        for col in range(3, len(technical_headers) + 1):
            ws3.cell(row=r, column=col, value=f"synthetic-tech-{col}")

    ws4 = wb.create_sheet("Résumé")
    ws4.cell(row=1, column=1, value="Résumé")
    ws4.cell(row=2, column=1, value="synthetic")

    ws5 = wb.create_sheet("Instructions")
    ws5.cell(row=1, column=1, value="Instructions synthétiques")

    sheets_to_keep = [ws1, ws2, ws3, ws4, ws5][:sheet_count]
    for extra in [ws1, ws2, ws3, ws4, ws5][sheet_count:]:
        wb.remove(extra)

    wb.save(str(path))
    return path


def make_default_rows(n=6):
    # 2 HIGH (OUI, recent), 2 MEDIUM (OUI, old), 1 EXCLUDED (NON), 1 SKIPPED_UNCERTAIN (INCERTAIN)
    return [
        {"decision": "OUI", "annee": 2021},
        {"decision": "OUI", "annee": 2024},
        {"decision": "OUI", "annee": 2010},
        {"decision": "OUI", "annee": 1999},
        {"decision": "NON", "annee": 2015},
        {"decision": "INCERTAIN", "annee": 2018},
    ][:n]


class WorkbookTestCase(unittest.TestCase):
    def setUp(self):
        self._tmpdir = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmpdir.cleanup)
        self.path = Path(self._tmpdir.name) / "synthetic_review.xlsx"

    def build(self, rows=None, **kwargs):
        rows = make_default_rows() if rows is None else rows
        build_synthetic_workbook(self.path, rows, **kwargs)
        return preview.sha256_file(self.path)


class TestHashVerification(WorkbookTestCase):
    def test_correct_hash_passes(self):
        digest = self.build()
        preview.verify_source_hash(self.path, digest)  # must not raise

    def test_incorrect_hash_raises(self):
        self.build()
        with self.assertRaises(preview.WorkbookValidationError):
            preview.verify_source_hash(self.path, "0" * 64)

    def test_incorrect_hash_is_case_insensitive_but_still_rejects_wrong_value(self):
        digest = self.build()
        preview.verify_source_hash(self.path, digest.upper())  # case-insensitive match is fine
        with self.assertRaises(preview.WorkbookValidationError):
            preview.verify_source_hash(self.path, digest[:-1] + ("0" if digest[-1] != "0" else "1"))


class TestValidWorkbook(WorkbookTestCase):
    def test_full_dry_run_passes_on_a_valid_workbook(self):
        digest = self.build()
        results = preview.run_dry_run(self.path, digest)
        as_dict = dict(results)
        self.assertEqual(as_dict["FILE_HASH_MATCH"], "PASS")
        self.assertEqual(as_dict["WORKBOOK_STRUCTURE"], "PASS")
        self.assertEqual(as_dict["IDENTIFIERS_UNIQUE"], "PASS")
        self.assertEqual(as_dict["ROW_ALIGNMENT"], "PASS")
        self.assertEqual(as_dict["DECISION_VOCABULARY"], "PASS")
        self.assertEqual(as_dict["YEAR_VALIDATION"], "PASS")
        self.assertEqual(as_dict["COUNT_INVARIANT"], "PASS")
        self.assertEqual(as_dict["DATABASE_ACCESSED"], "NO")
        self.assertEqual(as_dict["SOURCE_MODIFIED"], "NO")
        self.assertEqual(as_dict["DRY_RUN_RESULT"], "PASS")
        self.assertEqual(as_dict["TOTAL"], 6)


class TestStructuralValidation(WorkbookTestCase):
    def test_missing_sheet_fails_structure(self):
        digest = self.build(sheet_count=4)
        results = dict(preview.run_dry_run(self.path, digest))
        self.assertEqual(results["WORKBOOK_STRUCTURE"], "FAIL")
        self.assertEqual(results["DRY_RUN_RESULT"], "FAIL")

    def test_incorrect_header_fails_structure(self):
        bad_headers = ["Priorité", "WrongHeaderName", "Col3", "Col4", "Col5", "CDC utilisable ?", "Col7"]
        digest = self.build(main_headers=bad_headers)
        results = dict(preview.run_dry_run(self.path, digest))
        self.assertEqual(results["WORKBOOK_STRUCTURE"], "FAIL")
        self.assertEqual(results["DRY_RUN_RESULT"], "FAIL")

    def test_missing_decision_header_fails_structure(self):
        bad_headers = ["Priorité", "Année", "Col3", "Col4", "Col5", "SomethingElse", "Col7"]
        digest = self.build(main_headers=bad_headers)
        results = dict(preview.run_dry_run(self.path, digest))
        self.assertEqual(results["WORKBOOK_STRUCTURE"], "FAIL")


class TestIdentifierValidation(WorkbookTestCase):
    def test_duplicate_identifier_fails(self):
        digest = self.build(duplicate_primary_id_at=1)
        results = dict(preview.run_dry_run(self.path, digest))
        self.assertEqual(results["WORKBOOK_STRUCTURE"], "PASS")
        self.assertEqual(results["IDENTIFIERS_UNIQUE"], "FAIL")
        self.assertEqual(results["DRY_RUN_RESULT"], "FAIL")

    def test_missing_identifier_fails(self):
        digest = self.build(blank_primary_id_at=2)
        results = dict(preview.run_dry_run(self.path, digest))
        self.assertEqual(results["IDENTIFIERS_UNIQUE"], "FAIL")
        self.assertEqual(results["DRY_RUN_RESULT"], "FAIL")


class TestRowAlignment(WorkbookTestCase):
    def test_broken_row_alignment_fails(self):
        digest = self.build(technical_row_count=len(make_default_rows()) - 1)
        results = dict(preview.run_dry_run(self.path, digest))
        self.assertEqual(results["IDENTIFIERS_UNIQUE"], "PASS")
        self.assertEqual(results["ROW_ALIGNMENT"], "FAIL")
        self.assertEqual(results["DRY_RUN_RESULT"], "FAIL")

    def test_expected_total_mismatch_fails_row_alignment(self):
        digest = self.build()
        expected = preview.ExpectedCounts(total=999)
        results = dict(preview.run_dry_run(self.path, digest, expected))
        self.assertEqual(results["ROW_ALIGNMENT"], "FAIL")


class TestDecisionNormalization(WorkbookTestCase):
    def test_oui_variants_normalize_to_oui_and_route_by_year(self):
        rows = [
            {"decision": "oui", "annee": 2022},
            {"decision": " OUI ", "annee": 2022},
            {"decision": "Oui", "annee": 2010},
        ]
        digest = self.build(rows=rows)
        results = dict(preview.run_dry_run(self.path, digest))
        self.assertEqual(results["DECISION_VOCABULARY"], "PASS")
        self.assertEqual(results["HIGH"], 2)
        self.assertEqual(results["MEDIUM"], 1)
        self.assertEqual(results["EXCLUDED"], 0)
        self.assertEqual(results["SKIPPED_UNCERTAIN"], 0)

    def test_non_variants_normalize_to_non_and_are_excluded_regardless_of_year(self):
        rows = [
            {"decision": "non", "annee": 2022},
            {"decision": " NON ", "annee": 2005},
        ]
        digest = self.build(rows=rows)
        results = dict(preview.run_dry_run(self.path, digest))
        self.assertEqual(results["EXCLUDED"], 2)
        self.assertEqual(results["HIGH"], 0)
        self.assertEqual(results["MEDIUM"], 0)

    def test_incertain_routes_to_skipped_uncertain_not_excluded_not_usable(self):
        rows = [{"decision": "incertain", "annee": 2022}, {"decision": "Incertain", "annee": 2005}]
        digest = self.build(rows=rows)
        results = dict(preview.run_dry_run(self.path, digest))
        self.assertEqual(results["SKIPPED_UNCERTAIN"], 2)
        self.assertEqual(results["EXCLUDED"], 0)
        self.assertEqual(results["HIGH"], 0)
        self.assertEqual(results["MEDIUM"], 0)
        self.assertEqual(results["USABLE_TOTAL"], 0)

    def test_unknown_decision_value_is_rejected(self):
        rows = [{"decision": "PEUT-ETRE", "annee": 2022}]
        digest = self.build(rows=rows)
        results = dict(preview.run_dry_run(self.path, digest))
        self.assertEqual(results["DECISION_VOCABULARY"], "FAIL")
        self.assertEqual(results["DRY_RUN_RESULT"], "FAIL")


class TestYearValidation(WorkbookTestCase):
    def test_missing_year_fails(self):
        rows = [{"decision": "OUI", "annee": None}]
        digest = self.build(rows=rows)
        results = dict(preview.run_dry_run(self.path, digest))
        self.assertEqual(results["YEAR_VALIDATION"], "FAIL")
        self.assertEqual(results["DRY_RUN_RESULT"], "FAIL")

    def test_malformed_year_fails(self):
        for bad_year in ("abcd", "20x9", "99", "20211"):
            with self.subTest(bad_year=bad_year):
                path = Path(self._tmpdir.name) / f"malformed_{bad_year}.xlsx"
                build_synthetic_workbook(path, [{"decision": "OUI", "annee": bad_year}])
                digest = preview.sha256_file(path)
                results = dict(preview.run_dry_run(path, digest))
                self.assertEqual(results["YEAR_VALIDATION"], "FAIL")


class TestClassification(WorkbookTestCase):
    def test_high_calculation(self):
        rows = [{"decision": "OUI", "annee": y} for y in (2020, 2023, 2026)]
        digest = self.build(rows=rows)
        results = dict(preview.run_dry_run(self.path, digest))
        self.assertEqual(results["HIGH"], 3)
        self.assertEqual(results["MEDIUM"], 0)

    def test_medium_calculation(self):
        rows = [{"decision": "OUI", "annee": y} for y in (1980, 2000, 2019)]
        digest = self.build(rows=rows)
        results = dict(preview.run_dry_run(self.path, digest))
        self.assertEqual(results["MEDIUM"], 3)
        self.assertEqual(results["HIGH"], 0)

    def test_excluded_calculation(self):
        rows = [{"decision": "NON", "annee": y} for y in (2021, 1990)]
        digest = self.build(rows=rows)
        results = dict(preview.run_dry_run(self.path, digest))
        self.assertEqual(results["EXCLUDED"], 2)

    def test_uncertain_rows_are_preserved_and_skipped_not_counted_as_usable_or_excluded(self):
        rows = make_default_rows()  # includes exactly 1 INCERTAIN row
        digest = self.build(rows=rows)
        results = dict(preview.run_dry_run(self.path, digest))
        total = results["TOTAL"]
        usable = results["USABLE_TOTAL"]
        excluded = results["EXCLUDED"]
        skipped = results["SKIPPED_UNCERTAIN"]
        self.assertEqual(skipped, 1)
        # invariant: every row is accounted for exactly once
        self.assertEqual(usable + excluded + skipped, total)

    def test_count_invariant_holds_for_default_fixture(self):
        digest = self.build()
        results = dict(preview.run_dry_run(self.path, digest))
        self.assertEqual(
            results["HIGH"] + results["MEDIUM"] + results["EXCLUDED"] + results["SKIPPED_UNCERTAIN"],
            results["TOTAL"],
        )
        self.assertEqual(results["COUNT_INVARIANT"], "PASS")

    def test_count_invariant_failure_when_expected_counts_do_not_match(self):
        digest = self.build()  # 2 HIGH, 2 MEDIUM, 1 EXCLUDED, 1 SKIPPED_UNCERTAIN
        expected = preview.ExpectedCounts(total=6, high=999)
        results = dict(preview.run_dry_run(self.path, digest, expected))
        self.assertEqual(results["ROW_ALIGNMENT"], "PASS")
        self.assertEqual(results["COUNT_INVARIANT"], "FAIL")
        self.assertEqual(results["DRY_RUN_RESULT"], "FAIL")


class TestConfidentialityAndSafety(WorkbookTestCase):
    def test_no_row_level_content_in_printed_output(self):
        digest = self.build()
        results = preview.run_dry_run(self.path, digest)
        buffer = io.StringIO()
        with redirect_stdout(buffer):
            preview._print_results(results)
        output = buffer.getvalue()
        # None of the synthetic identifiers, decisions, or year values this
        # fixture used may appear verbatim in the printed output.
        for forbidden in ("SYN-ID-", "ALT-ID-", "synthetic-3", "synthetic-tech", "2021", "2024", "1999", "2010"):
            self.assertNotIn(forbidden, output)
        # only fixed labels and integer counts are allowed on each line
        for line in output.strip().splitlines():
            key, _, value = line.partition("=")
            self.assertTrue(key.isupper() or "_" in key)
            self.assertTrue(value in ("PASS", "FAIL", "YES", "NO") or value.lstrip("-").isdigit())

    def test_no_workbook_save_is_ever_called_during_a_dry_run(self):
        digest = self.build()
        with patch.object(openpyxl.Workbook, "save") as mock_save:
            preview.run_dry_run(self.path, digest)
            mock_save.assert_not_called()

    def test_module_never_references_database_or_credentials(self):
        # The module's own docstring *mentions* DATABASE_URL descriptively
        # (documenting that it never reads it) - so this checks actual usage
        # patterns, not a bare substring match against the whole file
        # (which would false-positive on that correct, honest documentation).
        source_text = MODULE_PATH.read_text(encoding="utf-8")
        for forbidden in ("psycopg", "pg8000", "import pg\n", "from pg ", "Pool(", "TEST_DATABASE_URL"):
            self.assertNotIn(forbidden, source_text)
        # No env-var access at all: the module never imports `os`, so it is
        # structurally incapable of reading DATABASE_URL or anything else
        # from the environment.
        self.assertNotIn("import os", source_text)
        self.assertNotIn("os.environ", source_text)
        self.assertNotIn("os.getenv", source_text)
        self.assertFalse(any(n == "os" for n in dir(preview)))

    def test_no_database_driver_module_is_imported_by_this_module(self):
        for name in list(sys.modules):
            self.assertFalse(
                name.startswith("psycopg") or name == "pg8000",
                f"unexpected database driver module present: {name}",
            )

    def test_cli_entrypoint_never_prints_row_level_data_on_success(self):
        digest = self.build()
        buffer = io.StringIO()
        err = io.StringIO()
        with redirect_stdout(buffer), redirect_stderr(err):
            exit_code = preview.main(
                ["--source", str(self.path), "--expected-sha256", digest, "--dry-run"]
            )
        self.assertEqual(exit_code, 0)
        output = buffer.getvalue()
        self.assertNotIn("SYN-ID-", output)
        self.assertIn("DRY_RUN_RESULT=PASS", output)

    def test_cli_entrypoint_fails_closed_on_wrong_hash(self):
        self.build()
        buffer = io.StringIO()
        with redirect_stdout(buffer):
            exit_code = preview.main(["--source", str(self.path), "--expected-sha256", "0" * 64, "--dry-run"])
        self.assertNotEqual(exit_code, 0)
        self.assertIn("FILE_HASH_MATCH=FAIL", buffer.getvalue())
        self.assertIn("DRY_RUN_RESULT=FAIL", buffer.getvalue())


if __name__ == "__main__":
    unittest.main()
